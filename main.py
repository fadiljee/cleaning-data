"""
Entrypoint CLI Utama untuk Pembersihan dan Validasi Data LK
"""
import argparse
import sys
import warnings
from pathlib import Path

# Abaikan peringatan format sel tanggal bawaan openpyxl dari file Excel BPS
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

import openpyxl

from src.cleaner import DataCleaner
from src.cleaner_lapangan import DataCleanerLapangan
from src.excel_exporter import ExcelExporter
from src.pdf_reporter import PDFReporter
from src.config import BASE_DIR, DEFAULT_INPUT_FILE, INPUT_DIR, OUTPUT_DIR


def find_input_file(custom_path: str = None) -> Path:
    """Mencari berkas input Excel di beberapa lokasi standar dengan prioritas data/input/"""
    if custom_path:
        p = Path(custom_path)
        if p.exists():
            return p
        raise FileNotFoundError(f"Berkas input tidak ditemukan: {custom_path}")

    # 1. Cek berkas spesifik di data/input/
    search_paths = [
        INPUT_DIR / DEFAULT_INPUT_FILE,
        INPUT_DIR / "LK.xlsx",
    ]
    for p in search_paths:
        if p.exists():
            return p

    # 2. Cari file .xlsx apapun yang valid di data/input/
    if INPUT_DIR.exists():
        for candidate in sorted(INPUT_DIR.glob("*.xlsx")):
            if not candidate.name.startswith("~") and not candidate.name.startswith("."):
                return candidate

    # 3. Fallback jika belum dipindah dari root project
    search_root = [
        BASE_DIR / DEFAULT_INPUT_FILE,
    ]
    for p in search_root:
        if p.exists():
            return p

    for candidate in BASE_DIR.glob("LK Cleaning Data*.xlsx"):
        if not candidate.name.startswith("~") and not candidate.name.startswith("."):
            return candidate

    raise FileNotFoundError(f"Tidak dapat menemukan berkas input di folder '{INPUT_DIR}' maupun '{BASE_DIR}'.")


def list_available_petugas(workbook_path: Path):
    """Menampilkan daftar seluruh nama petugas yang terdapat pada sheet LK"""
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    if "LK" not in wb.sheetnames:
        print("Sheet 'LK' tidak ditemukan.")
        return

    sheet = wb["LK"]
    petugas_counts = {}
    for row in sheet.iter_rows(min_row=3, values_only=True):
        p = row[30]
        if p:
            p_str = str(p).strip()
            petugas_counts[p_str] = petugas_counts.get(p_str, 0) + 1

    print("\n=== DAFTAR PETUGAS PADA SHEET LK ===")
    print(f"{'No':<4} | {'Nama Petugas':<45} | {'Jumlah Usaha':<12}")
    print("-" * 66)
    for idx, (p, cnt) in enumerate(sorted(petugas_counts.items(), key=lambda x: x[0]), start=1):
        print(f"{idx:<4} | {p:<45} | {cnt:<12}")
    print("-" * 66)
    print(f"Total Petugas: {len(petugas_counts)}")


def process_petugas(cleaner: DataCleaner, petugas_name: str, output_dir: Path):
    """Menjalankan alur lengkap pembersihan dan pembuatan berkas untuk satu petugas"""
    clean_name = petugas_name.replace(" ", "_").replace(".", "").replace(",", "")
    safe_name = "".join(c for c in clean_name if c.isalnum() or c == "_")

    print(f"\n[INFO] Membaca data untuk petugas: '{petugas_name}'...")
    raw_records = cleaner.load_records_by_petugas(petugas_name)

    if not raw_records:
        print(f"[PERINGATAN] Tidak ditemukan data untuk petugas '{petugas_name}'.")
        return

    print(f"[INFO] Ditemukan {len(raw_records)} usaha. Menjalankan algoritma pembersihan data...")
    cleaned_records, summary = cleaner.clean_records(raw_records)

    # Nama berkas output
    out_xlsx = output_dir / f"LK_Cleaning_Data_{safe_name}.xlsx"
    out_pdf = output_dir / f"Laporan_Cleaning_Data_{safe_name}.pdf"

    print(f"[INFO] Membuat berkas Excel: {out_xlsx.name}...")
    exporter = ExcelExporter(str(out_xlsx))
    exporter.export(cleaned_records, summary)

    print(f"[INFO] Membuat laporan PDF: {out_pdf.name}...")
    reporter = PDFReporter(str(out_pdf))
    reporter.generate(cleaned_records, summary)

    pos_count = summary.total_usaha - summary.total_selisih_minus
    pos_pct = (pos_count / summary.total_usaha) * 100 if summary.total_usaha > 0 else 0
    warna_pct = (summary.total_warna_hijau / summary.total_usaha) * 100 if summary.total_usaha > 0 else 0
    rasio_pct = (summary.total_rasio_hijau / summary.total_usaha) * 100 if summary.total_usaha > 0 else 0
    unique_ratios = len(set(round(c.rasio_ntb_output, 4) for c in cleaned_records))

    print("\n=== HASIL EVALUASI AKHIR ===")
    print(f"Petugas               : {summary.petugas}")
    print(f"Total Usaha           : {summary.total_usaha}")
    print(f"Status Warna Hijau    : {summary.total_warna_hijau}/{summary.total_usaha} ({warna_pct:.1f}% HIJAU, rentang {summary.min_cogs_ratio*100:.2f}% s.d. {summary.max_cogs_ratio*100:.2f}%)")
    print(f"Status Rasio NTB      : {summary.total_rasio_hijau}/{summary.total_usaha} ({rasio_pct:.1f}% HIJAU DALAM ACUAN, {unique_ratios} variasi rasio unik)")
    print(f"Selisih NTB >= 0      : {pos_count}/{summary.total_usaha} ({pos_pct:.1f}% untung sedikit/nol)")
    if summary.total_selisih_minus > 0:
        print(f"Penyesuaian Anomali   : {summary.total_selisih_minus} usaha (normalisasi output wajar dari anomali data mentah)")
    print(f"Berkas Excel          : {out_xlsx.resolve()}")
    print(f"Berkas PDF            : {out_pdf.resolve()}")


def main():
    parser = argparse.ArgumentParser(
        description="Sistem Otomasi Data Cleaning dan Validasi NTB Kategori G"
    )
    parser.add_argument(
        "--petugas",
        type=str,
        help="Nama petugas yang akan dibersihkan datanya (misal: 'Fadil Julianto' atau 'Emilia Prananda')"
    )
    parser.add_argument(
        "--list-petugas",
        action="store_true",
        help="Menampilkan seluruh daftar nama petugas yang ada di file Excel"
    )
    parser.add_argument(
        "--input",
        type=str,
        help="Jalur khusus ke berkas Excel input (opsional)"
    )
    parser.add_argument(
        "--output-dir",
        type=str,
        default=str(OUTPUT_DIR),
        help="Direktori penyimpanan hasil output Excel dan PDF (default: data/output)"
    )
    parser.add_argument(
        "--mode",
        type=str,
        choices=["lapangan", "standard"],
        default="lapangan",
        help="Mode pembersihan: 'lapangan' (sesuai lapangan, max 2x) atau 'standard'"
    )
    parser.add_argument(
        "--target-ntb",
        type=float,
        default=0.72,
        help="Target rasio NTB / Output (default: 0.72)"
    )
    parser.add_argument(
        "--rentang-bawah",
        type=float,
        default=0.64,
        help="Batas bawah rasio acuan NTB (default: 0.64)"
    )
    parser.add_argument(
        "--rentang-atas",
        type=float,
        default=0.78,
        help="Batas atas rasio acuan NTB (default: 0.78)"
    )
    parser.add_argument(
        "--min-surplus",
        type=float,
        default=0.0,
        help="Minimal surplus selisih NTB dalam Rupiah (default: 0.0, misal: 500000)"
    )
    parser.add_argument(
        "--max-multiplier",
        type=float,
        default=2.0,
        help="Batas maksimal kelipatan omzet dari nominal awal (default: 2.0, misal: 1.5)"
    )

    args = parser.parse_args()

    try:
        input_file = find_input_file(args.input)
    except FileNotFoundError as e:
        print(f"[ERROR] {e}", file=sys.stderr)
        sys.exit(1)

    if args.list_petugas:
        list_available_petugas(input_file)
        sys.exit(0)

    if not args.petugas:
        print("[INFO] Parameter --petugas tidak diberikan. Menjalankan default: 'Fadil Julianto'")
        petugas_target = "Fadil Julianto"
    else:
        petugas_target = args.petugas

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    if args.mode == "lapangan":
        print(
            f"[INFO] Menggunakan Mode 'Sesuai Lapangan' | Target NTB: {args.target_ntb:.2f} "
            f"| Rentang: [{args.rentang_bawah:.2f}, {args.rentang_atas:.2f}] "
            f"| Min Surplus: Rp {args.min_surplus:,.0f} | Max Multiplier: {args.max_multiplier:.1f}x"
        )
        cleaner = DataCleanerLapangan(
            str(input_file),
            target_ntb_ratio=args.target_ntb,
            min_ntb_ratio=args.rentang_bawah,
            max_ntb_ratio=args.rentang_atas,
            max_multiplier=args.max_multiplier,
            min_surplus=args.min_surplus,
        )
    else:
        print("[INFO] Menggunakan Mode 'Standard'")
        cleaner = DataCleaner(str(input_file))

    process_petugas(cleaner, petugas_target, output_dir)


if __name__ == "__main__":
    main()
